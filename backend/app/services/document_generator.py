"""In-memory document generation helpers."""

from io import BytesIO
from typing import Any


def _rows(content: str, data: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if data:
        return data
    return [{"content": content}] if content else []


def generate_pdf(content: str, data: list[dict[str, Any]]) -> BytesIO:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=LETTER, rightMargin=0.7 * inch, leftMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    story = []
    if content:
        for paragraph in content.splitlines() or [content]:
            story.append(Paragraph(paragraph.replace("&", "&amp;"), styles["BodyText"]))
            story.append(Spacer(1, 6))
    rows = _rows("", data)
    if rows:
        columns = list(rows[0].keys())
        table_data = [columns] + [[str(row.get(column, "")) for column in columns] for row in rows]
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f6654")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c8d8d0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    document.build(story or [Paragraph("", styles["BodyText"])])
    output.seek(0)
    return output


def generate_csv(content: str, data: list[dict[str, Any]]) -> BytesIO:
    import pandas as pd

    output = BytesIO()
    pd.DataFrame(_rows(content, data)).to_csv(output, index=False, encoding="utf-8-sig")
    output.seek(0)
    return output


def generate_xlsx(content: str, data: list[dict[str, Any]]) -> BytesIO:
    import pandas as pd

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(_rows(content, data)).to_excel(writer, index=False, sheet_name="Data")
    output.seek(0)
    return output


def generate_docx(content: str, data: list[dict[str, Any]]) -> BytesIO:
    from docx import Document

    output = BytesIO()
    document = Document()
    for paragraph in content.splitlines() if content else []:
        document.add_paragraph(paragraph)
    rows = _rows("", data)
    if rows:
        columns = list(rows[0].keys())
        table = document.add_table(rows=1, cols=len(columns))
        for cell, column in zip(table.rows[0].cells, columns):
            cell.text = str(column)
        for row in rows:
            cells = table.add_row().cells
            for cell, column in zip(cells, columns):
                cell.text = str(row.get(column, ""))
    if not content and not rows:
        document.add_paragraph("")
    document.save(output)
    output.seek(0)
    return output


GENERATORS = {
    "pdf": (generate_pdf, "application/pdf", ".pdf"),
    "csv": (generate_csv, "text/csv; charset=utf-8", ".csv"),
    "xlsx": (generate_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
    "docx": (generate_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
}
