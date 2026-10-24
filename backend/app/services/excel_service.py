"""Excel service for converting JSON tables to styled .xlsx downloads."""

from typing import List, Dict, Any, Optional


class ExcelService:
    """Converts JSON table data into professionally styled Excel workbook downloads."""

    @staticmethod
    def export(
        json_data: List[Dict[str, Any]],
        filename: str = "export.xlsx",
        sheet_name: str = "Data",
    ) -> Dict[str, Any]:
        """Export JSON table data to a styled Excel workbook.

        Args:
            json_data: List of dictionaries representing table rows.
            filename: Output filename for the Excel file.
            sheet_name: Name of the worksheet.

        Returns:
            Dict with success status and file path or error information.
        """
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError as e:
            return {
                "error": str(e),
                "filename": filename,
                "missing_packages": ["pandas", "openpyxl"],
            }

        try:
            df = pd.DataFrame(json_data)

            # Create workbook and add data
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name

            # Define styles
            header_font = Font(bold=True, size=11)
            header_fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
            header_font_white = Font(bold=True, size=11, color="FFFFFF")
            alignment_center = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # Write headers with styling
            if len(df.columns) > 0:
                for col_idx, column in enumerate(df.columns, 1):
                    cell = ws.cell(row=1, column=col_idx, value=column)
                    cell.font = header_font_white
                    cell.fill = header_fill
                    cell.alignment = alignment_center
                    cell.border = thin_border

            # Write data rows with styling
            for row_idx, (_, row) in enumerate(df.iterrows(), 2):
                for col_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.alignment = alignment_center
                    cell.border = thin_border

            # Adjust column widths automatically
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Save the workbook
            output_path = f"/tmp/{filename}"
            wb.save(output_path)

            return {
                "success": True,
                "filename": filename,
                "output_path": output_path,
                "rows": len(df),
                "columns": len(df.columns),
                "message": f"Successfully exported {len(df)} rows × {len(df.columns)} columns to Excel",
            }

        except Exception as e:
            return {"error": str(e), "filename": filename}